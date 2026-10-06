#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from io import BytesIO
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import ScatterChart, Reference, Series

LADO = {"ADELANTE": "Derecha", "ATRAS": "Izquierda", "PARAR": "Parado"}

# Datos: A #, B Fecha, C Hora, D Tiempo, E Sentido, F RPM, G Setpoint, H Error,
#        I PWM, J P, K I, L D, M Kp, N Ki, O Kd
COLS_DATOS = [
    ("#", 7, "0"),
    ("Fecha", 12, "dd/mm/yyyy"),
    ("Hora", 14, "hh:mm:ss.000"),
    ("Tiempo (s)", 11, "0.0"),
    ("Sentido", 11, "@"),
    ("Revoluciones (rpm)", 14, "0.0"),
    ("Setpoint (rpm)", 13, "0.0"),
    ("Error (rpm)", 11, "0.0"),
    ("PWM (%)", 10, "0.0"),
    ("P", 9, "0.00"),
    ("I", 9, "0.00"),
    ("D", 9, "0.00"),
    ("Kp", 8, "0.000"),
    ("Ki", 8, "0.000"),
    ("Kd", 8, "0.000"),
]

# Escalones: A n, B Hora, C Sentido, D Setpoint, E Inicio, F Duracion, G RPM ini, H RPM max,
#            I RPM fin, J Sobretiro, K Sobretiro %, L Error fin, M T estab, N Kp, O Ki, P Kd
COLS_ESC = [
    ("Escalón", 9, "0"),
    ("Hora inicio", 14, "hh:mm:ss"),
    ("Sentido", 11, "@"),
    ("Setpoint (rpm)", 13, "0.0"),
    ("Inicio (s)", 10, "0.0"),
    ("Duración (s)", 11, "0.0"),
    ("RPM inicial", 11, "0.0"),
    ("RPM máxima", 11, "0.0"),
    ("RPM final", 10, "0.0"),
    ("Sobretiro (rpm)", 13, "0.0"),
    ("Sobretiro (%)", 12, "0.0%"),
    ("Error final (rpm)", 13, "0.00"),
    ("T. establec. ±5% (s)", 15, "0.0"),
    ("Kp", 8, "0.000"),
    ("Ki", 8, "0.000"),
    ("Kd", 8, "0.000"),
]

letra = Font(name="Arial", size=10)
negrita = Font(name="Arial", size=10, bold=True)
titulo = Font(name="Arial", size=14, bold=True)
nota = Font(name="Arial", size=9, italic=True, color="666666")
encab = Font(name="Arial", size=10, bold=True, color="FFFFFF")
fondo_encab = PatternFill("solid", fgColor="1F2A44")
fondo_seccion = PatternFill("solid", fgColor="D9E2F3")
fondo_par = PatternFill("solid", fgColor="F2F5FA")
linea = Side(style="thin", color="B4BCCC")
borde = Border(left=linea, right=linea, top=linea, bottom=linea)


def num(x):
    try:
        return float(x)
    except ValueError:
        return x


def fecha_hora(txt):
    for fmt in ("%Y-%m-%d %H:%M:%S.%f", "%H:%M:%S.%f"):
        try:
            return datetime.strptime(txt, fmt)
        except ValueError:
            pass
    return None


def leer(lineas):
    # t, fecha hora, dir, sp, rpm, err, pwm, p, i, d, kp, ki, kd
    filas = []
    for lin in lineas:
        p = lin.split(",")
        if len(p) == 13:
            filas.append([num(p[0]), fecha_hora(p[1]), p[2]] + [num(x) for x in p[3:]])
    return filas


def encabezado(ws, cols, fila=1):
    for c, (nombre, ancho, _) in enumerate(cols, start=1):
        celda = ws.cell(row=fila, column=c, value=nombre)
        celda.font = encab
        celda.fill = fondo_encab
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = borde
        ws.column_dimensions[celda.column_letter].width = ancho
    ws.row_dimensions[fila].height = 32


def poner_fila(ws, fila, valores, cols):
    for c, v in enumerate(valores, start=1):
        celda = ws.cell(row=fila, column=c, value=v)
        celda.font = letra
        celda.number_format = cols[c - 1][2]
        celda.border = borde
        if fila % 2 == 0:
            celda.fill = fondo_par
        if cols[c - 1][2] in ("@", "dd/mm/yyyy", "hh:mm:ss.000", "hh:mm:ss"):
            celda.alignment = Alignment(horizontal="center")


def hoja_datos(ws, filas):
    encabezado(ws, COLS_DATOS)
    for n, f in enumerate(filas, start=1):
        t, cuando, sentido, sp, rpm, err, pwm, p, i, d, kp, ki, kd = f
        fecha = cuando.date() if cuando and cuando.year > 1900 else None
        hora = cuando.time() if cuando else None
        valores = [n, fecha, hora, t, LADO.get(sentido, sentido), rpm, sp, err, pwm, p, i, d, kp, ki, kd]
        poner_fila(ws, n + 1, valores, COLS_DATOS)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:O{len(filas) + 1}"


def tramos(filas):
    # un escalon = mismo sentido, setpoint y constantes
    clave = lambda f: (f[2], f[3], f[10], f[11], f[12])
    res = []
    ini = 0
    for k in range(1, len(filas) + 1):
        if k == len(filas) or clave(filas[k]) != clave(filas[ini]):
            if filas[ini][2] != "PARAR" and k - ini >= 3:
                res.append((ini + 2, k + 1))   # filas de excel en la hoja Datos
            ini = k
    return res


def hoja_escalones(ws, filas):
    ws["A1"] = "Análisis por escalón"
    ws["A1"].font = titulo
    ws["A2"] = ("Un escalón es cada tramo con el mismo sentido, setpoint y constantes. "
                "El sobretiro aplica cuando el setpoint sube o cambia el sentido (se mide desde que pasa por cero).")
    ws["A2"].font = nota

    encabezado(ws, COLS_ESC, fila=4)
    lista = tramos(filas)

    for n, (a, b) in enumerate(lista, start=1):
        r = n + 4
        ult10 = max(a, b - 9)
        valores = [
            n,
            f"=Datos!C{a}",
            f"=Datos!E{a}",
            f"=Datos!G{a}",
            f"=Datos!D{a}",
            f"=Datos!D{b}-Datos!D{a}",
            f"=Datos!F{a}",
            f"=MAX(Datos!F{a}:F{b})",
            f"=Datos!F{b}",
            (f'=IF(OR(D{r}>G{r},Datos!E{a}<>Datos!E{a - 1}),'
             f'MAX(0,MAX(INDEX(Datos!F{a}:F{b},MATCH(MIN(Datos!F{a}:F{b}),Datos!F{a}:F{b},0)):Datos!F{b})-D{r}),"-")'),
            f'=IF(ISNUMBER(J{r}),J{r}/D{r},"-")',
            f"=AVERAGE(Datos!H{ult10}:H{b})",
            (f'=IF(ABS(Datos!H{b})>0.05*D{r},"no se estabilizó",'
             f'IFERROR(LOOKUP(2,1/(ABS(Datos!H{a}:H{b})>0.05*D{r}),Datos!D{a + 1}:D{b + 1})-E{r},0))'),
            f"=Datos!M{b}",
            f"=Datos!N{b}",
            f"=Datos!O{b}",
        ]
        poner_fila(ws, r, valores, COLS_ESC)
        for c in (10, 11, 13):
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="right")

    ws.freeze_panes = "A5"
    if not lista:
        ws["A5"] = "No hubo escalones de al menos 3 muestras."
        ws["A5"].font = nota
    return len(lista)


def hoja_resumen(ws, est, ult, n_esc):
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 46

    ws["A1"] = "Control PID de motor con encoder"
    ws["A1"].font = titulo

    t = f"Datos!$D$2:$D${ult}"
    sen = f"Datos!$E$2:$E${ult}"
    rpm = f"Datos!$F$2:$F${ult}"
    sp = f"Datos!$G$2:$G${ult}"
    err = f"Datos!$H$2:$H${ult}"
    pwm = f"Datos!$I$2:$I${ult}"
    fin_esc = max(n_esc, 1) + 4
    esc = f"Escalones!$J$5:$J${fin_esc}"
    escp = f"Escalones!$K$5:$K${fin_esc}"

    secciones = [
        ("Prueba", [
            ("Fecha", "=Datos!B2", "dd/mm/yyyy", ""),
            ("Hora de inicio", "=Datos!C2", "hh:mm:ss", ""),
            ("Hora final", f"=Datos!C{ult}", "hh:mm:ss", ""),
            ("Exportado", datetime.now(), "dd/mm/yyyy hh:mm:ss", ""),
            ("Raspberry", "team5", "@", "L298N + encoder"),
        ]),
        ("Constantes PID al exportar", [
            ("Kp", est.get("kp", 0), "0.000", "proporcional"),
            ("Ki", est.get("ki", 0), "0.000", "integral"),
            ("Kd", est.get("kd", 0), "0.000", "derivativa"),
            ("Periodo de muestreo (s)",
             f"=IF(COUNT({t})>1,ROUND((MAX({t})-MIN({t}))/(COUNT({t})-1),3),0)",
             "0.000", "tiempo promedio entre filas de la hoja Datos"),
        ]),
        ("Revoluciones", [
            ("RPM máxima", f"=MAX({rpm})", "0.0", ""),
            ("RPM promedio", f"=AVERAGE({rpm})", "0.0", "incluye cuando va frenando"),
            ("Setpoint máximo (rpm)", f"=MAX({sp})", "0.0", ""),
            ("Error absoluto promedio (rpm)", f"=SUMPRODUCT(ABS({err}))/COUNT({err})", "0.00", "de todas las muestras"),
        ]),
        ("Sentido de giro", [
            ("Muestras a la derecha", f'=COUNTIF({sen},"Derecha")', "0", ""),
            ("Muestras a la izquierda", f'=COUNTIF({sen},"Izquierda")', "0", ""),
            ("Muestras parado / frenando", f'=COUNTIF({sen},"Parado")', "0", ""),
            ("Muestras totales", f"=COUNT({t})", "0", "una cada 0.1 s"),
            ("Duración (s)", f"=MAX({t})-MIN({t})", "0.0", ""),
        ]),
        ("PID", [
            ("PWM promedio (%)", f"=AVERAGE({pwm})", "0.0", ""),
            ("PWM máximo (%)", f"=MAX({pwm})", "0.0", ""),
            ("Escalones analizados", n_esc, "0", "ver hoja Escalones"),
            ("Sobretiro máximo (rpm)", f"=IF(COUNT({esc})>0,MAX({esc}),0)", "0.0", "el mayor de los escalones de subida"),
            ("Sobretiro máximo (%)", f"=IF(COUNT({escp})>0,MAX({escp}),0)", "0.0%", ""),
        ]),
    ]

    fila = 3
    for nombre, renglones in secciones:
        for col in "ABC":
            ws[f"{col}{fila}"].fill = fondo_seccion
        ws[f"A{fila}"] = nombre
        ws[f"A{fila}"].font = negrita
        fila += 1
        for etiqueta, valor, formato, txt in renglones:
            ws[f"A{fila}"] = etiqueta
            ws[f"A{fila}"].font = letra
            ws[f"B{fila}"] = valor
            ws[f"B{fila}"].font = negrita
            ws[f"B{fila}"].number_format = formato
            ws[f"B{fila}"].alignment = Alignment(horizontal="right")
            ws[f"C{fila}"] = txt
            ws[f"C{fila}"].font = nota
            ws[f"A{fila}"].border = borde
            ws[f"B{fila}"].border = borde
            fila += 1
        fila += 1

    ws[f"A{fila}"] = "Hojas: Datos (una fila por muestra), Escalones (análisis de cada cambio) y Gráficas."
    ws[f"A{fila}"].font = nota


def grafica(ws_datos, ult, cols, nombre_y, titulo_g):
    ch = ScatterChart()
    ch.title = titulo_g
    ch.style = 2
    ch.height = 10
    ch.width = 26
    ch.x_axis.title = "Tiempo (s)"
    ch.y_axis.title = nombre_y
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.legend.position = "b"

    x = Reference(ws_datos, min_col=4, min_row=2, max_row=ult)
    for col, color, punteada in cols:
        y = Reference(ws_datos, min_col=col, min_row=1, max_row=ult)
        s = Series(y, x, title_from_data=True)
        s.marker.symbol = "none"
        s.smooth = False
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 19050
        if punteada:
            s.graphicalProperties.line.dashStyle = "dash"
        ch.series.append(s)
    return ch


def armar_excel(lineas, est):
    filas = leer(lineas)
    ult = len(filas) + 1

    wb = Workbook()
    resumen = wb.active
    resumen.title = "Resumen"
    datos = wb.create_sheet("Datos")
    escalones = wb.create_sheet("Escalones")
    graf = wb.create_sheet("Gráficas")

    hoja_datos(datos, filas)
    n_esc = hoja_escalones(escalones, filas)
    hoja_resumen(resumen, est, ult, n_esc)

    graf["A1"] = "Respuesta del motor con control PID"
    graf["A1"].font = titulo
    graf.add_chart(grafica(datos, ult, [(7, "C9A800", True), (6, "0096A8", False)],
                           "RPM", "Setpoint vs revoluciones medidas"), "A3")
    graf.add_chart(grafica(datos, ult, [(9, "D6336C", False)],
                           "PWM (%)", "Salida del PID (PWM)"), "A25")

    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida
